"""End-to-end: real .pptx file -> Flask route -> JSON, entirely against stub_provider
(zero network calls).
"""
import io
from pptx import Presentation
from pptx.util import Inches


def _make_sample_pptx() -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    box = s1.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Our CNN model achieved 96% accuracy on the test set."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_analyze_rejects_non_pptx_file(client):
    data = {"file": (io.BytesIO(b"not a real file"), "notes.txt")}
    resp = client.post("/api/v1/presentation/analyze", data=data, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert resp.get_json()["error"]["code"] == "UNSUPPORTED_FORMAT"


def test_analyze_rejects_empty_file(client):
    data = {"file": (io.BytesIO(b""), "empty.pptx")}
    resp = client.post("/api/v1/presentation/analyze", data=data, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert resp.get_json()["error"]["code"] == "INVALID_FILE"


def test_analyze_rejects_corrupted_file(client):
    data = {"file": (io.BytesIO(b"invalid corrupt content that is not a zip"), "corrupt.pptx")}
    resp = client.post("/api/v1/presentation/analyze", data=data, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert resp.get_json()["error"]["code"] == "CORRUPTED_FILE"


def test_analyze_returns_valid_result_shape(client):
    pptx_bytes = _make_sample_pptx()
    data = {"file": (io.BytesIO(pptx_bytes), "sample.pptx")}
    resp = client.post("/api/v1/presentation/analyze", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200
    body = resp.get_json()

    assert body["status"] in ("completed", "partial")
    assert "scores" in body
    assert "summary" in body
    assert "claims" in body
    assert "verifications" in body
    assert "completeness" in body
    assert "issues" in body
    assert "suggested_interview_questions" in body
    assert "score_evidence" in body
    assert "brief" in body
    assert set(body["score_evidence"]) == {"overall", "fact_accuracy", "evidence_coverage", "reliability"}
    assert body["brief"]

    assert body["completeness"]["slides_total"] == 1
    assert body["completeness"]["slides_processed"] == 1
    assert body["completeness"]["slides_failed"] == 0
    assert len(body["claims"]) >= 1
