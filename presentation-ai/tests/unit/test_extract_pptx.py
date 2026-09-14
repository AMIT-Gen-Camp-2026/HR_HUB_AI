from io import BytesIO

from pptx import Presentation
from pptx.util import Inches

from app.pipeline.extract_pptx import extract


def _two_slide_bytes() -> bytes:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    box = s1.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Slide one body"
    s2 = prs.slides.add_slide(blank)
    box2 = s2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box2.text_frame.text = "Slide two body"
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_slide_numbers_are_one_indexed_and_preserved():
    result = extract(_two_slide_bytes())
    assert result.slide_count == 2
    assert [s.slide_number for s in result.slides] == [1, 2]
    assert "Slide one body" in result.slides[0].elements[0].content
    assert "Slide two body" in result.slides[1].elements[0].content


def test_extracts_speaker_notes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Slide content"
    slide.notes_slide.notes_text_frame.text = "Speaker note explaining methodology"

    buf = BytesIO()
    prs.save(buf)
    result = extract(buf.getvalue())
    assert result.slide_count == 1
    types = [el.type for el in result.slides[0].elements]
    assert "text" in types
    assert "note" in types
    assert any("Speaker note explaining methodology" in el.content for el in result.slides[0].elements)
