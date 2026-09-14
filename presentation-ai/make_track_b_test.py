from pptx import Presentation

prs = Presentation()

s1 = prs.slides.add_slide(prs.slide_layouts[1])
s1.shapes.title.text = 'Model Performance'
s1.placeholders[1].text_frame.text = 'Our CNN model achieved 99.9% accuracy on the validation set.'

s2 = prs.slides.add_slide(prs.slide_layouts[1])
s2.shapes.title.text = 'Dataset'
s2.placeholders[1].text_frame.text = 'The dataset contains 50,000 labeled images collected over six months.'

prs.save('track_b_test.pptx')
print('created')
