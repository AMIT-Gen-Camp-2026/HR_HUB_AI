"""PPTX bytes -> structured per-slide content. DETERMINISTIC — no AI, no provider
call. Every returned element MUST keep its slide_number: this is what lets the final
report point HR to the exact slide with an issue.
"""
from __future__ import annotations

import logging
import struct
import zipfile
from dataclasses import dataclass, field
from io import BytesIO

from pptx import Presentation
from pptx.exc import PackageNotFoundError

logger = logging.getLogger(__name__)


@dataclass
class SlideElement:
    type: str  # "text" | "table" | "note" | "chart"
    content: str


@dataclass
class SlideContent:
    slide_number: int
    title: str | None
    elements: list[SlideElement] = field(default_factory=list)
    extraction_error: str | None = None


@dataclass
class ExtractedPresentation:
    slide_count: int
    slides: list[SlideContent]
    failed_slide_numbers: list[int] = field(default_factory=list)


class PptxPackageCorrupted(Exception):
    """Raised when the .pptx bytes can't be parsed because the underlying zip/OPC
    package structure is malformed or truncated."""


_CORRUPTION_EXCEPTIONS = (
    PackageNotFoundError,  # python-pptx: "this isn't a valid OPC package"
    zipfile.BadZipFile,    # bytes aren't a zip archive at all
    KeyError,              # a required member is missing from the archive
    EOFError,              # archive truncated mid-read
    ValueError,            # e.g. "negative seek value" from a corrupted local header
    struct.error,          # malformed binary structure while parsing zip internals
)


def extract(file_bytes: bytes) -> ExtractedPresentation:
    try:
        prs = Presentation(BytesIO(file_bytes))
    except _CORRUPTION_EXCEPTIONS as exc:
        raise PptxPackageCorrupted(
            f"The .pptx package structure is malformed or truncated: {exc}"
        ) from exc

    slides: list[SlideContent] = []
    failed_slides: list[int] = []

    for idx, slide in enumerate(prs.slides, start=1):
        title = None
        elements: list[SlideElement] = []
        slide_err = None

        try:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    text = shape.text_frame.text.strip()
                    if title is None and shape == slide.shapes.title:
                        title = text
                        continue
                    elements.append(SlideElement(type="text", content=text))

                if shape.has_table:
                    table = shape.table
                    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                    rendered = "\n".join(" | ".join(r) for r in rows)
                    elements.append(SlideElement(type="table", content=rendered))

                if shape.has_chart:
                    try:
                        chart = shape.chart
                        chart_title = (
                            chart.chart_title.text_frame.text.strip()
                            if chart.has_title and chart.chart_title
                            else ""
                        )
                        series_lines = []
                        for s in chart.series:
                            series_lines.append(f"{getattr(s, 'name', 'Series')}: {list(getattr(s, 'values', []))}")
                        rendered_chart = f"Chart: {chart_title}\n" + "\n".join(series_lines) if chart_title else "\n".join(series_lines)
                        if rendered_chart.strip():
                            elements.append(SlideElement(type="chart", content=rendered_chart.strip()))
                    except Exception as e:
                        logger.debug("Failed to extract chart details on slide %d: %s", idx, e)

            # Extract speaker notes if present
            if hasattr(slide, "has_notes_slide") and slide.has_notes_slide:
                try:
                    notes_slide = slide.notes_slide
                    if notes_slide and notes_slide.notes_text_frame:
                        notes_text = notes_slide.notes_text_frame.text.strip()
                        if notes_text:
                            elements.append(SlideElement(type="note", content=notes_text))
                except Exception as e:
                    logger.debug("Failed to extract notes on slide %d: %s", idx, e)

        except Exception as exc:
            logger.warning("Error extracting shapes from slide %d: %s", idx, exc, exc_info=True)
            slide_err = str(exc)
            failed_slides.append(idx)

        slides.append(SlideContent(slide_number=idx, title=title, elements=elements, extraction_error=slide_err))

    return ExtractedPresentation(slide_count=len(slides), slides=slides, failed_slide_numbers=failed_slides)