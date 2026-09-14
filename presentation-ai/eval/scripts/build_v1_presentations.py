"""Build the v1 eval .pptx files (deterministic, no AI). Run from repo root:

    python eval/scripts/build_v1_presentations.py
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

OUT_DIR = Path(__file__).resolve().parent.parent / "datasets" / "presentation-extraction" / "v1" / "presentations"


def _blank() -> tuple[Presentation, object]:
    prs = Presentation()
    return prs, prs.slide_layouts[6]


def _add_text(slide, text: str, top: float = 0.5) -> None:
    box = slide.shapes.add_textbox(Inches(0.4), Inches(top), Inches(9.0), Inches(4.5))
    box.text_frame.word_wrap = True
    box.text_frame.text = text


def _add_title_and_body(prs, layout, title: str, body: str) -> None:
    slide = prs.slides.add_slide(layout)
    tbox = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    tbox.text_frame.text = title
    _add_text(slide, body, top=1.0)


def _save(prs: Presentation, name: str) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / name
    buf = BytesIO()
    prs.save(buf)
    path.write_bytes(buf.getvalue())
    print(f"wrote {path}")


def build_01_en_ml_results() -> None:
    prs, layout = _blank()
    _add_title_and_body(prs, layout, "Welcome to our project", "Welcome to our project. Agenda: results, data, and stack.")
    _add_title_and_body(
        prs, layout, "Results",
        "Our CNN model achieved 96% accuracy on the test set.\n"
        "The dataset contains 50,000 images across 10 classes.",
    )
    _add_title_and_body(
        prs, layout, "Architecture",
        "The system uses a microservices architecture.\n"
        "We used ResNet-50 as the backbone.",
    )
    _save(prs, "01_en_ml_results.pptx")


def build_02_en_objective_tech() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "Technology facts",
        "PostgreSQL supports JSON columns.\n"
        "FastAPI is built on Starlette.\n"
        "Prolog is a low-level language.",
    )
    _add_title_and_body(
        prs, layout, "Thank you",
        "Thank you. Questions?",
    )
    _save(prs, "02_en_objective_tech.pptx")


def build_03_ar_performance() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "نتائج المشروع",
        "حقق نموذجنا دقة 94% على مجموعة الاختبار.\n"
        "قاعدة البيانات تحتوي على 10000 صورة.",
    )
    _add_title_and_body(
        prs, layout, "التقنيات",
        "استخدمنا PostgreSQL كقاعدة بيانات أساسية.",
    )
    _save(prs, "03_ar_performance.pptx")


def build_04_mixed_lang() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "Mixed stack",
        "Our model achieved 88% F1-score.\n"
        "استخدمنا خوارزمية SVM للتصنيف.\n"
        "Redis supports pub/sub messaging.",
    )
    _save(prs, "04_mixed_lang.pptx")


def build_05_implausible() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "Perfect scores",
        "Our CNN model achieved 99.9% accuracy.\n"
        "Precision is 100%.\n"
        "The dataset contains 200 images.",
    )
    _save(prs, "05_implausible.pptx")


def build_06_metrics_table() -> None:
    prs, layout = _blank()
    slide = prs.slides.add_slide(layout)
    tbox = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.5))
    tbox.text_frame.text = "Evaluation metrics"
    table = slide.shapes.add_table(4, 3, Inches(0.5), Inches(1.2), Inches(8.5), Inches(2.4)).table
    rows = [
        ["Metric", "Value", "Split"],
        ["Accuracy", "91%", "test"],
        ["Recall", "87%", "test"],
        ["Latency", "45 ms", "inference"],
    ]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    slide2 = prs.slides.add_slide(layout)
    _add_text(slide2, "We trained on ImageNet-scale data with 1,200,000 samples.", top=0.5)
    _save(prs, "06_metrics_table.pptx")


def build_07_filler_only() -> None:
    prs, layout = _blank()
    _add_title_and_body(prs, layout, "Welcome", "Welcome to our project")
    _add_title_and_body(prs, layout, "Agenda", "Agenda\nIntroduction\nDemo\nQ&A")
    _add_title_and_body(prs, layout, "Thanks", "Thank you\nQuestions?")
    _save(prs, "07_filler_only.pptx")


def build_08_arch_algo_biz() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "System design",
        "The backend follows a hexagonal architecture.\n"
        "We implemented gradient boosting with XGBoost.\n"
        "The product can process invoices in under two seconds.\n"
        "The solution reduces operational cost by 30%.",
    )
    _save(prs, "08_arch_algo_biz.pptx")


def build_09_dataset_details() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "Dataset",
        "The dataset contains 50,000 images.\n"
        "There are 10 classes in the training set.\n"
        "Training samples were collected from public Kaggle sources.",
    )
    _save(prs, "09_dataset_details.pptx")


def build_10_latency_r2() -> None:
    prs, layout = _blank()
    _add_title_and_body(
        prs, layout, "Regression results",
        "The model reports an R2 score of 0.42.\n"
        "Median inference latency is 120 ms.\n"
        "Python lists are faster than NumPy arrays for large numeric workloads.",
    )
    _save(prs, "10_latency_r2.pptx")


def build_11_speaker_notes_details() -> None:
    prs, layout = _blank()
    
    # Slide 1: Model overview with details in speaker note
    s1 = prs.slides.add_slide(layout)
    t1 = s1.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    t1.text_frame.text = "LLM Fine-Tuning Overview"
    _add_text(s1, "Our customer support automation relies on custom instruction fine-tuning.", top=1.0)
    s1.notes_slide.notes_text_frame.text = "We fine-tuned LLaMA-3-8B on 15,000 customer service dialogues."

    # Slide 2: Deployment hardware with latency in note
    s2 = prs.slides.add_slide(layout)
    t2 = s2.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    t2.text_frame.text = "Serving Infrastructure"
    _add_text(s2, "Quantized model deployed on local edge hardware.", top=1.0)
    s2.notes_slide.notes_text_frame.text = "Inference achieves 35 tokens per second on an NVIDIA RTX 4090 GPU."

    # Slide 3: Database stack with an objective claim in note
    s3 = prs.slides.add_slide(layout)
    t3 = s3.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    t3.text_frame.text = "Persistence Tier"
    _add_text(s3, "Document storage for conversation histories.", top=1.0)
    s3.notes_slide.notes_text_frame.text = "MongoDB does not support ACID transactions across multiple documents."

    _save(prs, "11_speaker_notes_details.pptx")


def build_12_borderline_rag_ar_en() -> None:
    prs, layout = _blank()

    # Slide 1: Arabic RAG intro
    s1 = prs.slides.add_slide(layout)
    t1 = s1.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    t1.text_frame.text = "نظام استرجاع المعلومات القانونية RAG"
    _add_text(s1, "تم بناء خط الأنابيب باستخدام LangChain و FAISS لقاعدة معرفية تضم 25000 مستند قانوني.", top=1.0)

    # Slide 2: Table with retrieval metrics (borderline performance)
    s2 = prs.slides.add_slide(layout)
    t2 = s2.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.5))
    t2.text_frame.text = "مقاييس الاسترجاع (Retrieval Evaluation)"
    table = s2.shapes.add_table(3, 3, Inches(0.5), Inches(1.2), Inches(8.5), Inches(1.8)).table
    rows = [
        ["Metric", "Value", "Benchmark"],
        ["Recall@5", "78%", "Internal legal QA"],
        ["MRR", "0.65", "Internal legal QA"],
    ]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val

    # Slide 3: Objective search claim
    s3 = prs.slides.add_slide(layout)
    t3 = s3.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    t3.text_frame.text = "Search Backend"
    _add_text(s3, "Elasticsearch uses BM25 as its default ranking algorithm.", top=1.0)

    # Slide 4: Overstated / implausible satisfaction claim
    s4 = prs.slides.add_slide(layout)
    t4 = s4.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    t4.text_frame.text = "User Experience"
    _add_text(s4, "The system improved user query satisfaction to 99.8% with no observed hallucinations.", top=1.0)

    _save(prs, "12_borderline_rag_ar_en.pptx")


def main() -> None:
    build_01_en_ml_results()
    build_02_en_objective_tech()
    build_03_ar_performance()
    build_04_mixed_lang()
    build_05_implausible()
    build_06_metrics_table()
    build_07_filler_only()
    build_08_arch_algo_biz()
    build_09_dataset_details()
    build_10_latency_r2()
    build_11_speaker_notes_details()
    build_12_borderline_rag_ar_en()


if __name__ == "__main__":
    main()
