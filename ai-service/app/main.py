"""
app/main.py

نقطة الدخول الرئيسية للـ Flask API.
بيربط الـ pipeline اللي شغال بالفعل (extraction -> prompt -> model -> json)
خلف endpoints بسيطة.
"""

import logging
import os
import re
import tempfile

from flask import Flask, jsonify, request
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address

from pydantic import ValidationError as PydanticValidationError

from config.settings import config
from app.pipeline.ranking import rank as compute_ranking
from app.pipeline.run import extract_raw_text, clean_and_query
from app.pipeline.video_transcription import (
    ASRModelUnavailable,
    ASRInferenceError,
    VideoDecodeError,
    transcribe_video,
)
from app.providers.hf_provider import ModelInferenceError
from app.schemas.cv import CVSchema, JobDescription
from app.security.file_validator import (
    FileValidationError,
    generate_safe_storage_name,
    validate_extension,
    validate_file_content,
    validate_video_content,
    validate_video_mime_type,
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = max(config.MAX_CONTENT_LENGTH, config.VIDEO_MAX_SIZE_BYTES)

os.makedirs(config.UPLOAD_FOLDER, exist_ok=True)

# rate limiting عام على مستوى الـ app + حد أخص على endpoint الاستخراج
limiter = Limiter(get_remote_address, app=app, default_limits=["30 per hour"])


def _safe_video_error(error: Exception, status_code: int, message: str):
    """Return detail only for local development, never revealing paths or credentials."""
    body = {"success": False, "error": message}
    detail = getattr(error, "detail", None)
    if config.FLASK_DEBUG and detail:
        body["detail"] = re.sub(r"[A-Za-z]:\\[^\s]+|/[^\s]+", "[path]", detail)
    return jsonify(body), status_code


@app.errorhandler(413)
def request_too_large(_error):
    return jsonify({"success": False, "error": "Uploaded file exceeds the server size limit."}), 413


@app.errorhandler(429)
def request_rate_limited(_error):
    return jsonify({"success": False, "error": "Request rate limit exceeded. Please try again later."}), 429


@app.route("/api/v1/health", methods=["GET"])
def health():
    return jsonify({"status": "ok"}), 200


@app.route("/api/v1/cv/extract", methods=["POST"])
@limiter.limit("10 per hour")
def extract_cv():
    if "file" not in request.files:
        return jsonify({"success": False, "error": "No 'file' field in form-data."}), 400

    file = request.files["file"]

    if file.filename == "":
        return jsonify({"success": False, "error": "No file selected."}), 400

    try:
        ext = validate_extension(file.filename, config.ALLOWED_EXTENSIONS)
    except FileValidationError as e:
        return jsonify({"success": False, "error": str(e)}), 400

    temp_name = generate_safe_storage_name(file.filename)
    temp_path = os.path.join(config.UPLOAD_FOLDER, temp_name)

    try:
        file.save(temp_path)

        if os.path.getsize(temp_path) > config.MAX_CONTENT_LENGTH:
            return jsonify({"success": False, "error": "CV file exceeds the 10 MB limit."}), 413

        # تحقق من المحتوى الحقيقي للملف (magic bytes)، مش بس امتداده
        try:
            validate_file_content(temp_path, ext)
        except FileValidationError as e:
            return jsonify({"success": False, "error": str(e)}), 400

        raw_text = extract_raw_text(temp_path, ext)
        if not raw_text or not raw_text.strip():
            return jsonify({"success": False, "error": "No extractable text found in file."}), 422

        # clean_and_query دلوقتي بيتحقق من كل محاولة (parse_and_validate جوه
        # app/pipeline/run.py). لو موديل رجّع JSON تالف أو مش مطابق للـ schema،
        # هيتحسب فشل وهننتقل للموديل اللي بعده في MODEL_CHAIN تلقائيًا.
        try:
            validated_cv = clean_and_query(raw_text)
        except ModelInferenceError as e:
            logger.error("Model inference/validation failed across full chain: %s", e)
            return jsonify(
                {"success": False, "error": "Model inference failed. Please try again."}
            ), 502

        return jsonify({"success": True, "cv": validated_cv.model_dump()}), 200

    except Exception as e:
        logger.exception("Unexpected error during extraction")
        return jsonify({"success": False, "error": "Internal server error."}), 500

    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@app.route("/api/v1/rank", methods=["POST"])
@limiter.limit("60 per hour")
def rank_candidate():
    """
    بتاخد {"candidate": <CVSchema JSON>, "job_description": <JobDescription JSON>}
    وترجّع RankingResult. الـ candidate عادة هو ناتج /api/v1/cv/extract زي
    ما هو (validated_cv.model_dump()) - مفيش استخراج إضافي هنا، الـ scoring
    كله deterministic (algorithm، مش LLM call) في app/pipeline/ranking.py.
    """
    if not config.RANKING_ENABLED:
        return jsonify({"success": False, "error": "Ranking is switched off."}), 200

    payload = request.get_json(silent=True)
    if payload is None or not isinstance(payload, dict):
        return jsonify({"success": False, "error": "Request body must be a JSON object."}), 400

    try:
        candidate = CVSchema(**payload.get("candidate", {}))
        job_description = JobDescription(**payload.get("job_description", {}))
    except PydanticValidationError as e:
        return jsonify({"success": False, "error": e.errors()}), 422

    try:
        result = compute_ranking(candidate, job_description)
    except Exception:
        logger.exception("Unexpected error during ranking")
        return jsonify({"success": False, "error": "Internal server error."}), 500

    return jsonify({"success": True, "result": result.model_dump()}), 200


@app.route("/api/v1/video/transcribe", methods=["POST"])
@limiter.limit("10 per hour")
def transcribe_uploaded_video():
    """Transcribe an uploaded MP4/MOV/MKV/WebM without retaining the video."""
    if "file" not in request.files:
        return jsonify({"success": False, "error": "No 'file' field in form-data."}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"success": False, "error": "No video selected."}), 400

    temp_path: str | None = None
    try:
        try:
            extension = validate_extension(file.filename, config.VIDEO_ALLOWED_EXTENSIONS)
            validate_video_mime_type(file.mimetype)
        except FileValidationError as error:
            return jsonify({"success": False, "error": str(error)}), 415

        with tempfile.NamedTemporaryFile(delete=False, suffix=extension) as temporary_file:
            temp_path = temporary_file.name
            file.save(temporary_file)

        size = os.path.getsize(temp_path)
        if size == 0:
            return jsonify({"success": False, "error": "The uploaded video is empty."}), 422
        if size > config.VIDEO_MAX_SIZE_BYTES:
            return jsonify({"success": False, "error": "Video exceeds the configured size limit."}), 413
        try:
            validate_video_content(temp_path, extension)
        except FileValidationError as error:
            return jsonify({"success": False, "error": str(error)}), 415

        result = transcribe_video(temp_path)
        return jsonify({"success": True, **result}), 200
    except ASRModelUnavailable as error:
        logger.error("Video transcription unavailable at ASR setup: %s", error, exc_info=True)
        return _safe_video_error(error, 503, str(error))
    except VideoDecodeError as error:
        return jsonify({"success": False, "error": str(error)}), 422
    except ASRInferenceError as error:
        logger.error("Video transcription failed at ASR inference: %s", error, exc_info=True)
        return _safe_video_error(error, 500, str(error))
    except Exception:
        logger.exception("Unexpected error during video transcription")
        return jsonify({"success": False, "error": "Video transcription could not be completed."}), 500
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


@app.route("/api/v1/cv/evaluate", methods=["POST"])
@limiter.limit("10 per hour")
def evaluate_cv():
    import json as _json
    logger.info("POST /api/v1/cv/evaluate called")

    GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")

    if "file" not in request.files:
        return jsonify({"success": False, "error": "No 'file' field in form-data."}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"success": False, "error": "No file selected."}), 400

    jd_raw = request.form.get("job_description")
    job_description = None
    if jd_raw:
        try:
            jd_data = _json.loads(jd_raw)
            job_description = JobDescription(**jd_data)
        except Exception as e:
            logger.warning("Invalid job_description: %s", e)

    temp_name = generate_safe_storage_name(file.filename)
    temp_path = os.path.join(config.UPLOAD_FOLDER, temp_name)

    try:
        file.save(temp_path)
        if os.path.getsize(temp_path) > config.MAX_CONTENT_LENGTH:
            return jsonify({"success": False, "error": "CV file exceeds the 10 MB limit."}), 413

        ext = validate_extension(file.filename, config.ALLOWED_EXTENSIONS)
        try:
            validate_file_content(temp_path, ext)
        except FileValidationError as e:
            return jsonify({"success": False, "error": str(e)}), 400

        raw_text = extract_raw_text(temp_path, ext)
        if not raw_text or not raw_text.strip():
            return jsonify({"success": False, "error": "No extractable text found in file."}), 422

        if GEMINI_API_KEY:
            import urllib.request
            cv_prompt = (
                "Extract structured CV data from the following text. "
                "Return ONLY a JSON object with keys: personal_info (name, email, phone, location), "
                "education (array of degree, institution, year), "
                "experience (array of job_title, company, duration, description), "
                "skills (array of strings), inferred_skills (array), "
                "certifications (array), languages (array), "
                "projects (array of name, description, technologies_mentioned).\n\n"
                "CV TEXT:\n" + raw_text[:6000]
            )

            gemini_url = "https://generativelanguage.googleapis.com/v1beta/models/gemini-3-flash-preview:generateContent?key=" + GEMINI_API_KEY
            payload = _json.dumps({
                "contents": [{"parts": [{"text": cv_prompt}]}],
                "generationConfig": {"temperature": 0.2, "maxOutputTokens": 4096, "thinkingConfig": {"thinkingBudget": 0}}
            }).encode("utf-8")

            req = urllib.request.Request(gemini_url, data=payload, headers={"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=120) as resp:
                gemini_result = _json.loads(resp.read().decode("utf-8"))

            candidate = gemini_result["candidates"][0]
            content = candidate.get("content", {})
            parts = content.get("parts", [])
            raw_cv_text = ""
            for p in parts:
                if "text" in p:
                    raw_cv_text += p["text"]
            raw_cv_text = raw_cv_text.strip()
            if raw_cv_text.startswith("```"):
                raw_cv_text = re.sub(r"^```(?:json)?\s*", "", raw_cv_text)
                raw_cv_text = re.sub(r"\s*```$", "", raw_cv_text)

            try:
                cv_data = _json.loads(raw_cv_text)
            except _json.JSONDecodeError:
                fixed = re.sub(r',\s*}', '}', raw_cv_text)
                fixed = re.sub(r',\s*]', ']', fixed)
                cv_data = _json.loads(fixed)
        else:
            try:
                validated_cv = clean_and_query(raw_text)
                cv_data = validated_cv.model_dump()
            except ModelInferenceError as e:
                logger.error("Model inference failed: %s", e)
                return jsonify({"success": False, "error": "Model inference failed."}), 502

        # Skip strict CVSchema validation for Gemini output — Gemini may use
        # slightly different field names (e.g. "duration" vs "start_date").
        # Build a clean CVSchema-compatible dict for ranking.
        clean_cv = {
            "personal_info": cv_data.get("personal_info", {}),
            "education": [],
            "experience": [],
            "projects": cv_data.get("projects", []),
            "skills": cv_data.get("skills", []),
            "inferred_skills": cv_data.get("inferred_skills", []),
            "certifications": cv_data.get("certifications", []),
            "languages": cv_data.get("languages", []),
        }
        for edu in cv_data.get("education", []):
            clean_cv["education"].append({
                "degree": edu.get("degree"),
                "institution": edu.get("institution"),
                "graduation_year": edu.get("graduation_year") or edu.get("year"),
            })
        for exp in cv_data.get("experience", []):
            clean_cv["experience"].append({
                "job_title": exp.get("job_title") or exp.get("title"),
                "company": exp.get("company"),
                "start_date": exp.get("start_date"),
                "end_date": exp.get("end_date") or exp.get("duration"),
            })
        validated_cv = CVSchema.model_validate(clean_cv)

        result = {"success": True, "cv": cv_data}
        result["extraction_status"] = "completed"
        result["extraction_metadata"] = {}

        if job_description:
            try:
                ranking = compute_ranking(validated_cv, job_description)
                ranking_data = ranking.model_dump()
                cv_skills_lower = [sk.lower() for sk in (validated_cv.skills or [])]
                ranking_data["matched_required_skills"] = [
                    s for s in (job_description.required_skills or [])
                    if s.lower() in cv_skills_lower
                ]
                ranking_data["missing_required_skills"] = [
                    s for s in (job_description.required_skills or [])
                    if s.lower() not in cv_skills_lower
                ]
                ranking_data["matched_preferred_skills"] = [
                    s for s in (job_description.nice_to_have_skills or [])
                    if s.lower() in cv_skills_lower
                ]
                ranking_data["missing_preferred_skills"] = [
                    s for s in (job_description.nice_to_have_skills or [])
                    if s.lower() not in cv_skills_lower
                ]
                ranking_data["scoring_version"] = "1.0"
                result["ranking"] = ranking_data
            except Exception as e:
                logger.warning("Ranking failed: %s", e)
                result["ranking"] = None
                result["extraction_status"] = "completed_with_ranking_error"
        else:
            result["ranking"] = None

        return jsonify(result), 200

    except Exception as e:
        logger.exception("Unexpected error during CV evaluate")
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)}), 500

    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@app.route("/api/v1/presentation/analyze", methods=["POST"])
@limiter.limit("10 per hour")
def analyze_presentation():
    import uuid
    import json

    GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
    if not GEMINI_API_KEY:
        return jsonify({"success": False, "error": "GEMINI_API_KEY not configured."}), 500

    if "file" not in request.files:
        return jsonify({"success": False, "error": "No 'file' field in form-data."}), 400

    file = request.files["file"]
    if not file.filename:
        return jsonify({"success": False, "error": "No file selected."}), 400

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in {".pptx", ".ppt", ".pdf"}:
        return jsonify({"success": False, "error": f"Unsupported file type: {ext}. Use .pptx, .ppt, or .pdf"}), 400

    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            file.save(tmp)
            temp_path = tmp.name

        slides_text = ""
        if ext in {".pptx", ".ppt"}:
            try:
                from pptx import Presentation
                prs = Presentation(temp_path)
                for i, slide in enumerate(prs.slides, 1):
                    slide_texts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                    if slide_texts:
                        slides_text += f"\n--- Slide {i} ---\n" + "\n".join(slide_texts)
            except Exception as e:
                logger.warning("pptx extraction failed: %s, falling back to raw", e)
                slides_text = f"[Could not extract text from {file.filename}]"
        elif ext == ".pdf":
            try:
                import pdfplumber
                with pdfplumber.open(temp_path) as pdf:
                    for i, page in enumerate(pdf.pages, 1):
                        text = page.extract_text()
                        if text:
                            slides_text += f"\n--- Page {i} ---\n" + text
            except Exception as e:
                logger.warning("pdf extraction failed: %s", e)
                slides_text = f"[Could not extract text from {file.filename}]"

        if not slides_text.strip():
            return jsonify({"success": False, "error": "No extractable text found in the presentation."}), 422

        analysis_prompt = f"""You are an expert presentation evaluator for a hiring process.
Analyze the following presentation slides and return a JSON object with these exact fields:

{{
  "scores": {{
    "content_quality": <1-10>,
    "clarity": <1-10>,
    "structure": <1-10>,
    "visual_design": <1-10>,
    "overall": <1-10>
  }},
  "summary": {{
    "strengths": ["<strength1>", "<strength2>"],
    "weaknesses": ["<weakness1>", "<weakness2>"],
    "key_observations": ["<observation1>", "<observation2>"]
  }},
  "brief": "<2-3 sentence overall assessment>",
  "score_evidence": {{
    "content_quality": "<reasoning>",
    "clarity": "<reasoning>",
    "structure": "<reasoning>",
    "visual_design": "<reasoning>"
  }},
  "claims": ["<claim1>", "<claim2>"],
  "verifications": ["<verification1>", "<verification2>"],
  "issues": ["<issue1>", "<issue2>"],
  "suggested_interview_questions": ["<question1>", "<question2>", "<question3>"],
  "completeness": {{
    "has_cover_slide": <true/false>,
    "has_conclusion": <true/false>,
    "has_data_evidence": <true/false>,
    "slide_count_assessment": "<assessment>"
  }}
}}

Return ONLY the JSON object, no markdown, no extra text.

PRESENTATION CONTENT:
{slides_text[:8000]}"""

        import urllib.request
        gemini_url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-3-flash-preview:generateContent?key={GEMINI_API_KEY}"
        payload = json.dumps({
            "contents": [{"parts": [{"text": analysis_prompt}]}],
            "generationConfig": {
                "temperature": 0.3,
                "maxOutputTokens": 8192,
                "thinkingConfig": {"thinkingBudget": 0}
            }
        }).encode("utf-8")

        req = urllib.request.Request(gemini_url, data=payload, headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=120) as resp:
            gemini_result = json.loads(resp.read().decode("utf-8"))

        candidate = gemini_result["candidates"][0]
        content = candidate.get("content", {})
        parts = content.get("parts", [])
        raw_text = ""
        for p in parts:
            if "text" in p:
                raw_text += p["text"]
        raw_text = raw_text.strip()
        if raw_text.startswith("```"):
            raw_text = re.sub(r"^```(?:json)?\s*", "", raw_text)
            raw_text = re.sub(r"\s*```$", "", raw_text)

        try:
            analysis = json.loads(raw_text)
        except json.JSONDecodeError:
            logger.warning("Gemini returned invalid JSON, attempting to fix")
            fixed = re.sub(r',\s*}', '}', raw_text)
            fixed = re.sub(r',\s*]', ']', fixed)
            analysis = json.loads(fixed)

        raw_scores = analysis.get("scores", {})
        overall_score = raw_scores.get("overall", 5)

        scores_json = {
            "overall": overall_score,
            "fact_accuracy": raw_scores.get("content_quality", overall_score),
            "verified_ratio": raw_scores.get("clarity", overall_score),
            "evidence_coverage": raw_scores.get("structure", overall_score),
            "claim_reliability": raw_scores.get("visual_design", overall_score),
        }

        raw_issues = analysis.get("issues", [])
        issues_json = []
        for i, issue_text in enumerate(raw_issues):
            if isinstance(issue_text, str):
                severity = "high" if i < 1 else "medium" if i < 3 else "low"
                issues_json.append({
                    "severity": severity,
                    "slide_number": i + 1,
                    "claim_text": issue_text,
                    "correction": None,
                })
            elif isinstance(issue_text, dict):
                issues_json.append(issue_text)

        raw_questions = analysis.get("suggested_interview_questions", [])
        questions_json = []
        for q in raw_questions:
            if isinstance(q, str):
                questions_json.append({"suggested_question": q})
            elif isinstance(q, dict):
                questions_json.append(q)

        claims = analysis.get("claims", [])
        verifications = analysis.get("verifications", [])
        summary_json = {
            "total_claims": len(claims),
            "supported": len([c for c in claims if isinstance(c, str)]),
            "contradicted": 0,
            "plausibility_flag": len(issues_json),
            "unclear": len(verifications),
        }

        return jsonify({
            "analysis_id": str(uuid.uuid4()),
            "status": "completed",
            "scores": scores_json,
            "summary": summary_json,
            "brief": analysis.get("brief", ""),
            "score_evidence": analysis.get("score_evidence", {}),
            "claims": claims,
            "verifications": verifications,
            "issues": issues_json,
            "suggested_interview_questions": questions_json,
            "completeness": analysis.get("completeness", {}),
        }), 200

    except Exception as e:
        logger.exception("Presentation analysis failed")
        return jsonify({"success": False, "error": str(e)}), 500

    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


@app.route("/api/v1/chat", methods=["POST"])
def chat():
    """AI Chatbot endpoint — accepts messages and returns a reply via Gemini."""
    import json
    import urllib.request

    GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
    if not GEMINI_API_KEY:
        return jsonify({"error": "GEMINI_API_KEY not configured"}), 500

    data = request.get_json(force=True, silent=True)
    if not data:
        return jsonify({"error": "No JSON payload provided"}), 400

    messages = data.get("messages", [])
    prompt = data.get("prompt", "")

    if not messages and not prompt:
        return jsonify({"error": "Either messages or prompt is required"}), 400

    # Build conversation history for Gemini
    gemini_contents = []
    for msg in messages:
        role = msg.get("role", "user")
        content = msg.get("content", "")
        if role == "system":
            continue
        gemini_contents.append({
            "role": "model" if role == "assistant" else "user",
            "parts": [{"text": content}]
        })

    # The last user message
    user_text = prompt or (messages[-1]["content"] if messages else "")

    system_instruction = (
        "You are an AI assistant for the AMIT HR Hub recruitment system. "
        "You help hiring managers, track heads, and reviewers with:\n"
        "- CV screening questions and candidate evaluation\n"
        "- Presentation analysis and interview preparation\n"
        "- Qualification pipeline stage guidance\n"
        "- Policy and process questions about recruitment\n"
        "- Data-driven insights about candidates and job profiles\n\n"
        "Be concise, professional, and helpful. If you don't know something, say so. "
        "Always answer in the same language the user writes in."
    )

    try:
        # Build Gemini API request
        system_msg = {"role": "user", "parts": [{"text": system_instruction}]}
        model_reply = {"role": "model", "parts": [{"text": "Understood. I will help with AMIT HR Hub questions."}]}

        contents = [system_msg, model_reply] + gemini_contents if gemini_contents else [
            system_msg, model_reply,
            {"role": "user", "parts": [{"text": user_text}]}
        ]

        # If only prompt (no history), just send it
        if not gemini_contents:
            contents = [
                {"role": "user", "parts": [{"text": user_text}]}
            ]

        payload = json.dumps({
            "contents": contents,
            "generationConfig": {
                "temperature": 0.7,
                "maxOutputTokens": 2048,
            }
        }).encode("utf-8")

        gemini_url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-3.1-flash-lite:generateContent?key={GEMINI_API_KEY}"
        req = urllib.request.Request(gemini_url, data=payload, headers={"Content-Type": "application/json"})

        with urllib.request.urlopen(req, timeout=30) as resp:
            gemini_result = json.loads(resp.read().decode("utf-8"))

        candidate = gemini_result["candidates"][0]
        parts = candidate.get("content", {}).get("parts", [])
        reply_text = "".join(p.get("text", "") for p in parts)

        return jsonify({
            "reply": reply_text,
            "model_version": "gemini-3.1-flash-lite",
            "tokens_in": sum(len(m.get("content", "").split()) for m in messages),
            "tokens_out": len(reply_text.split()),
        }), 200

    except Exception as e:
        logger.exception("Chat request failed")
        return jsonify({"error": str(e)}), 500


if __name__ == "__main__":
    config.validate()
    app.run(debug=config.FLASK_DEBUG, port=config.FLASK_PORT)
